#!/usr/bin/env python3
"""早會全勤王：讀取 Excel【記錄原檔】，依規則計算出席與獎勵。

用法：
  python3 tools/attendance_king.py 記錄表.xlsx config.json out.json

config.json（不放進 Git，存在 D1 attendance.meta 的 config 欄位）：
  {
    "start": "2026-10-02", "end": "2026-12-30",
    "weekly_from": "2026-10-05",          # 這天以前的週不發咖啡券（只算出席）
    "holidays": {"2026-10-09": "國定假日"}, # 放假日：全員算出席，不送週三加碼
    "leave": {"2026-10-16": ["姓名", ...]}  # 公假：算出席
  }

規則：
  算出席＝出席、公假、放假；遲到、早退、沒紀錄＝未計出席
  當週完全出席（該週每天都算出席）＝咖啡券 1 張
  週三實際出席或公假＝咖啡券加碼 1 張
  單月全勤＝當月每個早會日都算出席（星巴克禮券 2 張）
  全勤之王＝期間內每個早會日都算出席（餐券 2 張）
"""
import json, re, sys
from datetime import date, datetime, timedelta

import openpyxl

PRESENT = {"出席", "公假", "放假"}
CODE = {"出席": "O", "公假": "L", "放假": "H", "未計出席": "X", "未到": "F"}


def norm(s):
    return re.sub(r"[\s　·．.\-_()（）]", "", str(s or "")).lower()


def d2s(d):
    return d.isoformat()


def meeting_days(start, end):
    d, out = start, []
    while d <= end:
        if d.weekday() in (0, 2, 4):
            out.append(d)
        d += timedelta(days=1)
    return out


def read_records(path):
    """回傳 {日期: {姓名: 狀態}} 與【全勤王】名單（若有）。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["記錄原檔"]
    recs = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if not isinstance(v, (datetime, date)):
            continue
        day = v.date() if isinstance(v, datetime) else v
        names = {}
        for r in range(2, ws.max_row + 1):
            n = ws.cell(r, c + 1).value
            s = ws.cell(r, c + 2).value
            if n and str(n).strip():
                names[str(n).strip()] = str(s or "").strip()
        recs[day] = names
    roster = []
    if "全勤王" in wb.sheetnames:
        w2 = wb["全勤王"]
        for r in range(4, w2.max_row + 1):
            n = w2.cell(r, 3).value
            if n and str(n).strip() and str(n).strip() != "報名業務員":
                roster.append(str(n).strip())
    return recs, roster


def compute(recs, roster, cfg):
    start = date.fromisoformat(cfg["start"])
    end = date.fromisoformat(cfg["end"])
    weekly_from = date.fromisoformat(cfg.get("weekly_from", cfg["start"]))
    hol = {date.fromisoformat(k): v for k, v in cfg.get("holidays", {}).items()}
    leave = {date.fromisoformat(k): {norm(x) for x in v} for k, v in cfg.get("leave", {}).items()}
    days = meeting_days(start, end)
    last = max([d for d in recs if start <= d <= end], default=start - timedelta(days=1))
    resolved = lambda d: d <= last or d in hol

    # 姓名：記錄原檔＋全勤王名單，依第一次出現的寫法
    names, seen = [], set()
    for d in sorted(recs):
        for n in recs[d]:
            if norm(n) not in seen:
                seen.add(norm(n)); names.append(n)
    for n in roster:
        if norm(n) not in seen:
            seen.add(norm(n)); names.append(n)

    lookup = {d: {norm(k): v for k, v in m.items()} for d, m in recs.items()}

    def status(n, d):
        k = norm(n)
        if d in hol:
            return "放假"
        if not resolved(d):
            return "未到"
        if k in leave.get(d, set()):
            return "公假"
        s = lookup.get(d, {}).get(k, "")
        return "出席" if s == "出席" else ("未計出席" if s or d in lookup else "未計出席")

    # 週：依週一分組，週歸屬月份＝週一所在月份
    weeks = {}
    for d in days:
        mon = d - timedelta(days=d.weekday())
        weeks.setdefault(mon, []).append(d)

    week_list = [(mon, ds) for mon, ds in sorted(weeks.items())
                 if mon >= weekly_from - timedelta(days=weekly_from.weekday())]
    wk_meta = [{"from": d2s(ds[0]), "to": d2s(ds[-1]), "m": mon.month if mon >= start else start.month}
               for mon, ds in week_list]

    people = []
    for n in names:
        st = {d2s(d): status(n, d) for d in days}
        wk_code, coupons_by_month = "", {}
        for (mon, ds), wm in zip(week_list, wk_meta):
            done = all(resolved(d) for d in ds)
            full = done and all(st[d2s(d)] in PRESENT for d in ds)
            wed = [d for d in ds if d.weekday() == 2]
            bonus = bool(wed) and resolved(wed[0]) and st[d2s(wed[0])] in ("出席", "公假")
            coupons_by_month[wm["m"]] = coupons_by_month.get(wm["m"], 0) + int(full) + int(bonus)
            wk_code += str(int(done) * 4 + int(full) * 2 + int(bonus))  # 0-7：完成/全勤/週三加碼
        full_months = []
        for m in sorted({d.month for d in days}):
            mds = [d for d in days if d.month == m]
            if all(resolved(d) and st[d2s(d)] in PRESENT for d in mds):
                full_months.append(m)
        all_rs = [d for d in days if resolved(d)]
        king_alive = all(st[d2s(d)] in PRESENT for d in all_rs)
        king = king_alive and len(all_rs) == len(days)
        parts = [f"{m}月咖啡券 {c} 張" for m, c in sorted(coupons_by_month.items()) if c]
        parts += [f"{m}月全勤" for m in full_months]
        if king:
            parts.append("全勤之王")
        people.append({"name": n, "award": "・".join(parts), "coupons": sum(coupons_by_month.values()),
                       "days": "".join(CODE[st[d2s(d)]] for d in days),  # 對應 calendar 順序
                       "weeks": wk_code})                                # 對應 weeks 順序
    cal = [{"d": d2s(d), "type": "holiday" if d in hol else ("done" if resolved(d) else "future"),
            "note": hol.get(d, "")} for d in days]
    return {"lastRecord": d2s(last), "calendar": cal, "weeks": wk_meta, "people": people}


def main():
    xlsx, cfgp, outp = sys.argv[1:4]
    cfg = json.load(open(cfgp, encoding="utf-8"))
    recs, roster = read_records(xlsx)
    out = compute(recs, roster, cfg)
    json.dump(out, open(outp, "w", encoding="utf-8"), ensure_ascii=False)
    win = [p for p in out["people"] if p["award"]]
    print(f"紀錄到 {out['lastRecord']}；{len(out['people'])} 人；達成名單 {len(win)} 人")
    for p in sorted(win, key=lambda p: -p["coupons"]):
        print(f"  {p['name']}\t{p['award']}")


if __name__ == "__main__":
    main()
